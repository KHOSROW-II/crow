"""
Business logic: import, sync, query, export, CRUD, platform registry, flags,
dynamic data sources, per-contact attachment folders, and DB-stored images.
"""

import base64 as _b64
import csv
import hashlib as _hashlib
import json
import os as _os
import shutil
import sqlite3
import subprocess
import sys
from io import BytesIO as _BytesIO
from pathlib import Path

from db import (
    DB_PATH, INPUT_JSON,
    connect, normalize_phone, normalize_national_code, normalize_name,
    audit,
    list_platforms as db_list_platforms,
    platform_names as db_platform_names,
    add_platform as db_add_platform,
    rename_platform as db_rename_platform,
    remove_platform as db_remove_platform,
)
from data_sources import (
    list_sources as ds_list,
    get_source as ds_get,
    add_source as ds_add,
    update_source as ds_update,
    remove_source as ds_remove,
    source_label_map as ds_labels,
    source_db_names as ds_db_names,
    source_exists as ds_exists,
)


WORKDIR         = Path(_os.environ.get("CONTACTS_WORKDIR", "."))
ATTACHMENTS_DIR = WORKDIR / "attachments"


# ------------------------------------------------------------
# Logger
# ------------------------------------------------------------
_log_fn = print

def set_logger(fn):
    global _log_fn
    _log_fn = fn

def log(*args, **kwargs):
    _log_fn(*args, **kwargs)


# ------------------------------------------------------------
# Caches
# ------------------------------------------------------------
_platforms_cache = None
_source_labels_cache = None
_external_indexes = {}


def _invalidate_platforms():
    global _platforms_cache
    _platforms_cache = None


def _invalidate_source_labels():
    global _source_labels_cache
    _source_labels_cache = None


def invalidate_external_cache(name=None):
    global _external_indexes
    if name:
        _external_indexes.pop(name, None)
    else:
        _external_indexes = {}


# ------------------------------------------------------------
# Data source wrappers
# ------------------------------------------------------------
def list_data_sources(only_enabled=False):
    return ds_list(only_enabled=only_enabled)


def get_data_source(name):
    return ds_get(name)


def add_data_source(s):
    n = ds_add(s)
    _invalidate_source_labels()
    log(f"[source] added: {n}")
    return n


def update_data_source(name, s):
    n = ds_update(name, s)
    _invalidate_source_labels()
    log(f"[source] updated: {n}")
    return n


def remove_data_source(name):
    ds_remove(name)
    _invalidate_source_labels()
    invalidate_external_cache(name)
    log(f"[source] removed: {name}")


def source_label_map():
    global _source_labels_cache
    if _source_labels_cache is None:
        _source_labels_cache = ds_labels()
    return _source_labels_cache


def source_db_names():
    return ds_db_names()


# ------------------------------------------------------------
# Platform registry
# ------------------------------------------------------------
def list_platforms():
    global _platforms_cache
    if _platforms_cache is None:
        _platforms_cache = db_list_platforms()
    return _platforms_cache


def platform_names():
    return [p["name"] for p in list_platforms()]


def add_platform(name: str, display_name: str | None = None):
    n = db_add_platform(name, display_name)
    _invalidate_platforms()
    log(f"[platform] added: {n}")
    return n


def rename_platform(name: str, display_name: str | None = None):
    db_rename_platform(name, display_name)
    _invalidate_platforms()
    log(f"[platform] renamed display name of {name} -> {display_name}")


def remove_platform(name):
    db_remove_platform(name)
    _invalidate_platforms()
    log(f"[platform] removed: {name}")


# ------------------------------------------------------------
# Attachments (filesystem)
# ------------------------------------------------------------
def contact_folder(contact_id, create=True) -> Path | None:
    if contact_id is None:
        return None
    folder = ATTACHMENTS_DIR / str(contact_id)
    if create:
        folder.mkdir(parents=True, exist_ok=True)
    return folder


def copy_attachment(contact_id, source_path) -> str:
    src = Path(source_path).expanduser()
    if not src.exists():
        raise FileNotFoundError(str(src))
    src = src.resolve()

    folder = contact_folder(contact_id, create=True)
    if folder is None:
        raise ValueError("Cannot copy attachment without a contact.")

    dest = folder / src.name
    try:
        if dest.exists() and dest.resolve() == src:
            try:
                return str(dest.relative_to(Path.cwd()))
            except Exception:
                return str(dest)
    except Exception:
        pass

    n = 1
    while dest.exists():
        dest = folder / f"{src.stem}_{n}{src.suffix}"
        n += 1

    shutil.copy2(src, dest)
    try:
        return str(dest.relative_to(Path.cwd()))
    except Exception:
        return str(dest)


def set_contact_photo(contact_id, source_path) -> str:
    """Legacy: copy file to attachments/ and store the path."""
    if contact_id is None:
        raise ValueError("Save the contact before attaching a photo.")

    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT photo_path FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
    old = row["photo_path"] if row else None

    new_path = copy_attachment(contact_id, source_path)

    with connect(DB_PATH) as conn:
        conn.execute(
            "UPDATE contacts SET photo_path=?, updated_at=CURRENT_TIMESTAMP "
            "WHERE id=?",
            (new_path, contact_id),
        )
    audit("set_contact_photo", f"id={contact_id}, path={new_path}")

    if old and old != new_path:
        try:
            old_p = Path(old)
            if not old_p.is_absolute():
                old_p = Path.cwd() / old_p
            new_p = Path(new_path)
            if not new_p.is_absolute():
                new_p = Path.cwd() / new_p
            if old_p.exists() and old_p.resolve() != new_p.resolve():
                old_p.unlink()
        except Exception:
            pass

    return new_path


def clear_contact_photo(contact_id):
    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT photo_path FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
        if not row:
            return
        old = row["photo_path"]
        conn.execute(
            "UPDATE contacts SET photo_path=NULL, updated_at=CURRENT_TIMESTAMP "
            "WHERE id=?",
            (contact_id,),
        )
    if old:
        try:
            old_p = Path(old)
            if not old_p.is_absolute():
                old_p = Path.cwd() / old_p
            if old_p.exists():
                old_p.unlink()
        except Exception:
            pass
    audit("clear_contact_photo", f"id={contact_id}")


def photo_abs_path(photo_path) -> Path | None:
    if not photo_path:
        return None
    p = Path(photo_path)
    if not p.is_absolute():
        p = Path.cwd() / p
    return p if p.exists() else None


def open_contact_folder(contact_id):
    folder = contact_folder(contact_id, create=True)
    if folder is None:
        return
    folder = folder.resolve()
    try:
        if sys.platform == "win32":
            _os.startfile(str(folder))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(folder)])
        else:
            subprocess.Popen(["xdg-open", str(folder)])
    except Exception as e:
        log(f"[attach] could not open folder: {e}")


def list_contact_files(contact_id):
    folder = contact_folder(contact_id, create=False)
    if folder is None or not folder.exists():
        return []
    out = []
    for p in sorted(folder.iterdir()):
        if p.is_file():
            try:
                size = p.stat().st_size
            except Exception:
                size = 0
            out.append((p.name, str(p), size))
    return out


# ============================================================
# Contact images stored in the DB (raw bytes + base64)
# ============================================================
_MIME_BY_FORMAT = {
    "PNG":  "image/png",
    "JPG":  "image/jpeg",
    "JPEG": "image/jpeg",
    "GIF":  "image/gif",
    "BMP":  "image/bmp",
    "WEBP": "image/webp",
    "TIFF": "image/tiff",
    "TIF":  "image/tiff",
    "ICO":  "image/x-icon",
    "PPM":  "image/x-portable-pixmap",
    "PGM":  "image/x-portable-graymap",
}


def _detect_image_meta(data: bytes):
    """Return (mime, width, height). Best effort, never raises."""
    # 1) PIL for full info
    try:
        from PIL import Image as _PILImage  # type: ignore
        img = _PILImage.open(_BytesIO(data))
        w, h = img.size
        fmt = (img.format or "").upper()
        mime = _MIME_BY_FORMAT.get(fmt, "image/unknown")
        return mime, w, h
    except Exception:
        pass

    # 2) fallback: magic numbers only
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg", None, None
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png", None, None
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif", None, None
    if data[:2] == b"BM":
        return "image/bmp", None, None
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", None, None
    if data[:4] in (b"II*\x00", b"MM\x00*"):
        return "image/tiff", None, None
    return "application/octet-stream", None, None


def _insert_image(conn, contact_id, data: bytes, name: str, set_primary: bool):
    mime, w, h = _detect_image_meta(data)
    sha = _hashlib.sha256(data).hexdigest()
    b64 = _b64.b64encode(data).decode("ascii")

    if set_primary:
        conn.execute(
            "UPDATE contact_images SET is_primary=0 WHERE contact_id=?",
            (contact_id,),
        )

    cur = conn.execute("""
        INSERT INTO contact_images
            (contact_id, name, mime, size_bytes, width, height,
             sha256, raw_data, encoded_data, encoding, is_primary)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'base64', ?)
    """, (contact_id, name, mime, len(data), w, h, sha,
          sqlite3.Binary(data), b64, 1 if set_primary else 0))
    return cur.lastrowid


def add_contact_image(contact_id, file_path, *, name=None,
                      set_primary=False):
    """Read a file from disk and store it in contact_images."""
    if contact_id is None:
        raise ValueError("Save the contact before adding images.")
    p = Path(file_path).expanduser()
    if not p.exists():
        raise FileNotFoundError(str(p))
    data = p.read_bytes()
    if name is None:
        name = p.name
    with connect(DB_PATH) as conn:
        img_id = _insert_image(conn, contact_id, data, name, set_primary)
    audit("add_contact_image",
          f"cid={contact_id}, img={img_id}, name={name}")
    log(f"[images] added #{img_id} for contact #{contact_id}")
    return img_id


def add_contact_image_from_bytes(contact_id, data: bytes, *,
                                 name: str = "image",
                                 set_primary: bool = False):
    if contact_id is None:
        raise ValueError("Save the contact before adding images.")
    if not data:
        raise ValueError("Empty image data.")
    with connect(DB_PATH) as conn:
        img_id = _insert_image(conn, contact_id, data, name, set_primary)
    audit("add_contact_image",
          f"cid={contact_id}, img={img_id}, from_bytes")
    return img_id


def list_contact_images(contact_id):
    """Return metadata rows (no BLOB payloads)."""
    with connect(DB_PATH) as conn:
        return [dict(r) for r in conn.execute("""
            SELECT id, contact_id, name, mime, size_bytes, width, height,
                   sha256, is_primary, encoding, created_at
            FROM contact_images
            WHERE contact_id=?
            ORDER BY is_primary DESC, id ASC
        """, (contact_id,))]


def get_contact_image_bytes(image_id):
    """Return the raw bytes of a stored image."""
    if image_id is None:
        return None
    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT raw_data FROM contact_images WHERE id=?",
            (image_id,),
        ).fetchone()
    if not row or row["raw_data"] is None:
        return None
    return bytes(row["raw_data"])


def get_contact_image_encoded(image_id):
    """Return the base64-encoded text of a stored image."""
    if image_id is None:
        return None
    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT encoded_data FROM contact_images WHERE id=?",
            (image_id,),
        ).fetchone()
    if not row or row["encoded_data"] is None:
        return None
    return str(row["encoded_data"])


def get_primary_image_id(contact_id):
    if contact_id is None:
        return None
    with connect(DB_PATH) as conn:
        row = conn.execute("""
            SELECT id FROM contact_images
            WHERE contact_id=?
            ORDER BY is_primary DESC, id ASC
            LIMIT 1
        """, (contact_id,)).fetchone()
    return row["id"] if row else None


def set_primary_image(contact_id, image_id):
    with connect(DB_PATH) as conn:
        conn.execute(
            "UPDATE contact_images SET is_primary=0 WHERE contact_id=?",
            (contact_id,),
        )
        conn.execute(
            "UPDATE contact_images SET is_primary=1 "
            "WHERE id=? AND contact_id=?",
            (image_id, contact_id),
        )
    audit("set_primary_image", f"cid={contact_id}, img={image_id}")


def delete_contact_image(image_id):
    if image_id is None:
        return
    with connect(DB_PATH) as conn:
        conn.execute("DELETE FROM contact_images WHERE id=?", (image_id,))
    audit("delete_contact_image", f"img={image_id}")
    log(f"[images] deleted #{image_id}")


def delete_all_contact_images(contact_id):
    with connect(DB_PATH) as conn:
        conn.execute("DELETE FROM contact_images WHERE contact_id=?",
                     (contact_id,))
    audit("delete_all_contact_images", f"cid={contact_id}")


def export_contact_image(image_id, dest_path):
    data = get_contact_image_bytes(image_id)
    if data is None:
        return None
    Path(dest_path).write_bytes(data)
    audit("export_contact_image", f"img={image_id} -> {dest_path}")
    return str(dest_path)


def migrate_photo_paths_to_db():
    """
    One-shot migration: for every contact with a photo_path, copy the file
    into contact_images and mark it primary.
    """
    migrated = 0
    with connect(DB_PATH) as conn:
        rows = conn.execute(
            "SELECT id, photo_path FROM contacts "
            "WHERE photo_path IS NOT NULL AND photo_path != ''"
        ).fetchall()

    for row in rows:
        cid = row["id"]
        p = row["photo_path"]
        try:
            existing = get_primary_image_id(cid)
            if existing is not None:
                continue
            add_contact_image(cid, p, set_primary=True)
            migrated += 1
        except Exception as e:
            log(f"[migrate] contact #{cid} failed: {e}")

    log(f"[migrate] {migrated} photo_path(s) moved into contact_images.")
    return migrated


# ------------------------------------------------------------
# External index per source
# ------------------------------------------------------------
def _load_index(src):
    idx = {}
    path = src["db_path"]
    if not Path(path).exists():
        return idx
    try:
        conn = sqlite3.connect(path)
        conn.row_factory = sqlite3.Row
        table = src["table_name"]
        for r in conn.execute(f"SELECT * FROM {table}"):
            d = dict(r)
            for field in src["phone_fields"]:
                if field not in d:
                    continue
                phone = normalize_phone(d.get(field))
                if phone:
                    idx[phone] = d
        conn.close()
    except Exception as e:
        log(f"[source] index failed for {src['name']}: {e}")
    return idx


def _get_index(src):
    name = src["name"]
    if name in _external_indexes:
        return _external_indexes[name]
    idx = _load_index(src)
    _external_indexes[name] = idx
    return idx


# ------------------------------------------------------------
# Low-level
# ------------------------------------------------------------
VALID_STATUS = {"unchecked", "checking", "verified", "invalid", "anonymous"}


def _is_yes(v):
    if v is None:
        return False
    return str(v).strip().lower() in ("yes", "y", "true", "1", "بله")


def _upsert_contact(conn, phone, *, status=None, name=None, last_name=None,
                    national_code=None, notes=None, photo_path=None):
    phone = normalize_phone(phone)
    if not phone:
        return None
    n  = normalize_name(name)
    ln = normalize_name(last_name)

    row = conn.execute(
        "SELECT id FROM contacts WHERE phone_number=?", (phone,)
    ).fetchone()
    if row:
        conn.execute("""
            UPDATE contacts SET
                status         = COALESCE(?, status),
                name           = COALESCE(?, name),
                last_name      = COALESCE(?, last_name),
                name_norm      = COALESCE(?, name_norm),
                last_name_norm = COALESCE(?, last_name_norm),
                national_code  = COALESCE(?, national_code),
                notes          = COALESCE(?, notes),
                photo_path     = COALESCE(?, photo_path),
                updated_at     = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (status, name, last_name, n, ln, national_code, notes,
              photo_path, row["id"]))
        return row["id"]

    cur = conn.execute("""
        INSERT INTO contacts
            (phone_number, status, name, last_name,
             name_norm, last_name_norm, national_code, notes, photo_path)
        VALUES (?, COALESCE(?, 'unchecked'), ?, ?, ?, ?, ?, ?, ?)
    """, (phone, status, name, last_name, n, ln, national_code,
          notes, photo_path))
    return cur.lastrowid


def _set_social(conn, contact_id, platform, *,
                exists_status=None, username=None, display_name=None,
                user_id=None, verified_status=None, raw_data=None):
    if contact_id is None:
        return
    row = conn.execute(
        "SELECT id FROM contact_socials WHERE contact_id=? AND platform=?",
        (contact_id, platform)
    ).fetchone()
    if row:
        conn.execute("""
            UPDATE contact_socials SET
                exists_status   = COALESCE(?, exists_status),
                username        = COALESCE(?, username),
                display_name    = COALESCE(?, display_name),
                user_id         = COALESCE(?, user_id),
                verified_status = COALESCE(?, verified_status),
                raw_data        = COALESCE(?, raw_data),
                updated_at      = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (exists_status, username, display_name,
              user_id, verified_status, raw_data, row["id"]))
    else:
        conn.execute("""
            INSERT INTO contact_socials
                (contact_id, platform, exists_status, username,
                 display_name, user_id, verified_status, raw_data)
            VALUES (?, ?, COALESCE(?, 'unknown'), ?, ?, ?,
                    COALESCE(?, 'unchecked'), ?)
        """, (contact_id, platform, exists_status, username,
              display_name, user_id, verified_status, raw_data))


def _add_link(conn, contact_id, source_db, source_table, source_pk,
              match_key, match_value, confidence=1.0):
    conn.execute("""
        INSERT OR IGNORE INTO contact_links
            (contact_id, source_db, source_table, source_pk,
             match_key, match_value, confidence)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (contact_id, source_db, source_table, str(source_pk),
          match_key,
          str(match_value) if match_value is not None else None,
          confidence))


# ------------------------------------------------------------
# Public CRUD
# ------------------------------------------------------------
def create_contact(phone, *, status="unchecked", name=None, last_name=None,
                   national_code=None, notes=None, photo_path=None,
                   pinned=0, special=0):
    phone_norm = normalize_phone(phone)
    if not phone_norm:
        raise ValueError("Invalid phone number.")

    with connect(DB_PATH) as conn:
        existing = conn.execute(
            "SELECT id FROM contacts WHERE phone_number=?", (phone_norm,)
        ).fetchone()
        if existing:
            raise ValueError(
                f"Phone {phone_norm} already exists (contact #{existing['id']})."
            )
        cid = _upsert_contact(
            conn, phone_norm,
            status=status, name=name, last_name=last_name,
            national_code=national_code, notes=notes, photo_path=photo_path,
        )
        conn.execute(
            "UPDATE contacts SET pinned=?, special=? WHERE id=?",
            (1 if pinned else 0, 1 if special else 0, cid),
        )
    audit("create_contact", f"id={cid}, phone={phone_norm}")
    log(f"[edit] created contact #{cid} ({phone_norm})")
    return cid


def update_contact(contact_id, *, phone=None, status=None, name=None,
                   last_name=None, national_code=None,
                   notes=None, photo_path=None, pinned=None, special=None):
    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT * FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
        if not row:
            raise ValueError(f"Contact #{contact_id} not found.")

        new_phone = normalize_phone(phone) if phone else row["phone_number"]
        if not new_phone:
            raise ValueError("Invalid phone number.")

        if new_phone != row["phone_number"]:
            other = conn.execute(
                "SELECT id FROM contacts WHERE phone_number=? AND id<>?",
                (new_phone, contact_id)
            ).fetchone()
            if other:
                raise ValueError(
                    f"Phone {new_phone} already used by contact #{other['id']}."
                )

        if notes is None:
            new_notes = row["notes"]
        else:
            new_notes = str(notes).strip() or None

        if photo_path is None:
            new_photo = row["photo_path"]
        else:
            new_photo = str(photo_path).strip() or None

        conn.execute("""
            UPDATE contacts SET
                phone_number   = ?,
                status         = COALESCE(?, status),
                name           = ?,
                last_name      = ?,
                name_norm      = ?,
                last_name_norm = ?,
                national_code  = ?,
                notes          = ?,
                photo_path     = ?,
                pinned         = COALESCE(?, pinned),
                special        = COALESCE(?, special),
                updated_at     = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (
            new_phone, status,
            name, last_name,
            normalize_name(name), normalize_name(last_name),
            national_code,
            new_notes,
            new_photo,
            None if pinned  is None else (1 if pinned  else 0),
            None if special is None else (1 if special else 0),
            contact_id,
        ))
    audit("update_contact", f"id={contact_id}")
    log(f"[edit] updated contact #{contact_id}")


def set_notes(contact_id, notes):
    update_contact(contact_id, notes=notes or "")


def delete_contact(contact_id):
    """Delete contact, its attachment folder, and DB images."""
    folder = contact_folder(contact_id, create=False)
    if folder is not None and folder.exists():
        try:
            shutil.rmtree(folder, ignore_errors=True)
        except Exception:
            pass

    with connect(DB_PATH) as conn:
        conn.execute("DELETE FROM contact_images WHERE contact_id=?",
                     (contact_id,))
        conn.execute("DELETE FROM contacts WHERE id=?", (contact_id,))

    audit("delete_contact", f"id={contact_id}")
    log(f"[edit] deleted contact #{contact_id}")


def update_social(contact_id, platform, *, exists_status=None,
                  username=None, display_name=None, user_id=None,
                  verified_status=None):
    if platform not in platform_names():
        raise ValueError(f"Unknown platform: {platform}")
    with connect(DB_PATH) as conn:
        _set_social(
            conn, contact_id, platform,
            exists_status=exists_status,
            username=username,
            display_name=display_name,
            user_id=user_id,
            verified_status=verified_status,
        )
    audit("update_social", f"cid={contact_id}, platform={platform}")


def delete_social(contact_id, platform):
    with connect(DB_PATH) as conn:
        conn.execute(
            "DELETE FROM contact_socials WHERE contact_id=? AND platform=?",
            (contact_id, platform)
        )
    audit("delete_social", f"cid={contact_id}, platform={platform}")
    log(f"[edit] removed {platform} row for contact #{contact_id}")


# ------------------------------------------------------------
# Flags
# ------------------------------------------------------------
def toggle_pin(contact_id):
    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT pinned FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
        if not row:
            raise ValueError(f"Contact #{contact_id} not found.")
        new_val = 0 if row["pinned"] else 1
        conn.execute(
            "UPDATE contacts SET pinned=?, updated_at=CURRENT_TIMESTAMP "
            "WHERE id=?",
            (new_val, contact_id),
        )
    audit("toggle_pin", f"id={contact_id}, pinned={new_val}")
    log(f"[edit] contact #{contact_id} pinned = {bool(new_val)}")
    return bool(new_val)


def toggle_special(contact_id):
    with connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT special FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
        if not row:
            raise ValueError(f"Contact #{contact_id} not found.")
        new_val = 0 if row["special"] else 1
        conn.execute(
            "UPDATE contacts SET special=?, updated_at=CURRENT_TIMESTAMP "
            "WHERE id=?",
            (new_val, contact_id),
        )
    audit("toggle_special", f"id={contact_id}, special={new_val}")
    log(f"[edit] contact #{contact_id} special = {bool(new_val)}")
    return bool(new_val)


def set_pinned_many(ids, pinned):
    if not ids:
        return 0
    val = 1 if pinned else 0
    with connect(DB_PATH) as conn:
        placeholders = ",".join(["?"] * len(ids))
        conn.execute(
            f"UPDATE contacts SET pinned=?, updated_at=CURRENT_TIMESTAMP "
            f"WHERE id IN ({placeholders})",
            [val, *ids],
        )
    audit("set_pinned_many", f"count={len(ids)}, pinned={val}")
    log(f"[edit] pinned set to {bool(val)} for {len(ids)} contact(s)")
    return len(ids)


def set_special_many(ids, special):
    if not ids:
        return 0
    val = 1 if special else 0
    with connect(DB_PATH) as conn:
        placeholders = ",".join(["?"] * len(ids))
        conn.execute(
            f"UPDATE contacts SET special=?, updated_at=CURRENT_TIMESTAMP "
            f"WHERE id IN ({placeholders})",
            [val, *ids],
        )
    audit("set_special_many", f"count={len(ids)}, special={val}")
    log(f"[edit] special set to {bool(val)} for {len(ids)} contact(s)")
    return len(ids)


# ------------------------------------------------------------
# Import from JSON
# ------------------------------------------------------------
def import_from_json(json_path: str | None = None) -> int:
    if json_path is None:
        json_path = INPUT_JSON
    p = Path(json_path)
    if not p.exists():
        log(f"[import] JSON file not found: {json_path}")
        return 0

    with open(p, encoding="utf-8") as f:
        data = json.load(f)

    contacts = data.get("contacts", [])
    imported = 0
    plats = platform_names()

    with connect(DB_PATH) as conn:
        for row in contacts:
            if not isinstance(row, dict):
                continue
            phone = row.get("phone")
            status = str(row.get("status") or "unchecked").strip().lower()
            if status not in VALID_STATUS:
                status = "unchecked"

            cid = _upsert_contact(
                conn, phone, status=status,
                name=row.get("name"),
                last_name=row.get("last_name"),
                national_code=row.get("national_code"),
                notes=row.get("notes"),
            )
            if cid is None:
                continue
            imported += 1

            if "pinned" in row or "special" in row:
                conn.execute(
                    "UPDATE contacts SET "
                    "pinned  = COALESCE(?, pinned), "
                    "special = COALESCE(?, special) "
                    "WHERE id = ?",
                    (
                        None if "pinned"  not in row else (1 if row["pinned"]  else 0),
                        None if "special" not in row else (1 if row["special"] else 0),
                        cid,
                    ),
                )

            for platform in plats:
                if platform not in row:
                    continue
                _set_social(
                    conn, cid, platform,
                    exists_status="yes" if _is_yes(row.get(platform)) else "no",
                    username=row.get(f"{platform}_username"),
                    display_name=row.get(f"{platform}_display_name"),
                    user_id=row.get(f"{platform}_userid"),
                    verified_status=status,
                )

    log(f"[import] {imported} contacts imported from {json_path}.")
    audit("import_json", f"{json_path}, count={imported}")
    return imported


# ------------------------------------------------------------
# Matching
# ------------------------------------------------------------
def _find_contact_by_phone(conn, phone):
    phone = normalize_phone(phone)
    if not phone:
        return None
    row = conn.execute(
        "SELECT id FROM contacts WHERE phone_number=?", (phone,)
    ).fetchone()
    return row["id"] if row else None


def _find_contact_by_national_code(conn, national_code):
    nc = normalize_national_code(national_code)
    if not nc:
        return None
    row = conn.execute(
        "SELECT id FROM contacts WHERE national_code=?", (nc,)
    ).fetchone()
    return row["id"] if row else None


def _find_contacts_by_name(conn, name, last_name):
    n, f = normalize_name(name), normalize_name(last_name)
    if not n or not f:
        return []
    rows = conn.execute(
        "SELECT id FROM contacts WHERE name_norm=? AND last_name_norm=?",
        (n, f)
    ).fetchall()
    return [r["id"] for r in rows]


def _update_contact_names(conn, cid, first, last):
    conn.execute("""
        UPDATE contacts SET
            name           = COALESCE(name, ?),
            last_name      = COALESCE(last_name, ?),
            name_norm      = COALESCE(name_norm, ?),
            last_name_norm = COALESCE(last_name_norm, ?),
            updated_at     = CURRENT_TIMESTAMP
        WHERE id = ?
    """, (first, last, normalize_name(first), normalize_name(last), cid))


def _update_contact_national_code(conn, cid, nc):
    nc = normalize_national_code(nc)
    if not nc:
        return
    conn.execute("""
        UPDATE contacts SET
            national_code = COALESCE(national_code, ?),
            updated_at    = CURRENT_TIMESTAMP
        WHERE id = ?
    """, (nc, cid))


def match_with_source(src) -> int:
    if not Path(src["db_path"]).exists():
        log(f"[sync] source '{src['name']}' db not found: {src['db_path']}")
        return 0

    matched = 0
    stats = {}

    try:
        src_conn = sqlite3.connect(src["db_path"])
        src_conn.row_factory = sqlite3.Row
    except Exception as e:
        log(f"[sync] cannot open source '{src['name']}': {e}")
        return 0

    try:
        with connect(DB_PATH) as conn:
            real_cols = {r["name"] for r in
                         src_conn.execute(f"PRAGMA table_info({src['table_name']})")}
            phone_fields = [f for f in (src.get("phone_fields") or [])
                            if f in real_cols]
            name_cols = src.get("name_columns") or []
            first_col = name_cols[0] if len(name_cols) > 0 and name_cols[0] in real_cols else None
            last_col  = name_cols[1] if len(name_cols) > 1 and name_cols[1] in real_cols else None
            nc_col = src.get("national_code_column")
            if nc_col and nc_col not in real_cols:
                nc_col = None

            for r in src_conn.execute(f"SELECT * FROM {src['table_name']}"):
                cid = None
                match_key = None
                match_value = None
                confidence = 0.0

                if "phone" in src["matchers"]:
                    for field in phone_fields:
                        phone = normalize_phone(r[field])
                        if not phone:
                            continue
                        c = _find_contact_by_phone(conn, phone)
                        if c:
                            cid, match_key = c, f"phone:{field}"
                            match_value, confidence = phone, 1.0
                            break

                if cid is None and "national_code" in src["matchers"] and nc_col:
                    nc = r[nc_col]
                    if nc and not normalize_phone(nc):
                        c = _find_contact_by_national_code(conn, nc)
                        if c:
                            cid, match_key = c, "national_code"
                            match_value = normalize_national_code(nc)
                            confidence = 0.95

                if cid is None and "name_family" in src["matchers"]:
                    n = r[first_col] if first_col else None
                    f = r[last_col] if last_col else None
                    if n and f:
                        ids = _find_contacts_by_name(conn, n, f)
                        if len(ids) == 1:
                            cid, match_key = ids[0], "name_family"
                            match_value = f"{normalize_name(n)} {normalize_name(f)}"
                            confidence = 0.7
                        elif len(ids) > 1:
                            cid, match_key = ids[0], "name_family_ambiguous"
                            match_value = f"{normalize_name(n)} {normalize_name(f)}"
                            confidence = 0.4

                if cid is None:
                    if src.get("create_missing") and phone_fields:
                        phone = None
                        for field in phone_fields:
                            phone = normalize_phone(r[field])
                            if phone:
                                break
                        if phone:
                            n = r[first_col] if first_col else None
                            f = r[last_col] if last_col else None
                            cur = conn.execute("""
                                INSERT INTO contacts
                                    (phone_number, status, name, last_name,
                                     name_norm, last_name_norm)
                                VALUES (?, 'unchecked', ?, ?, ?, ?)
                            """, (phone, n, f, normalize_name(n), normalize_name(f)))
                            cid = cur.lastrowid
                            match_key = f"created_from_{src['name']}"
                            match_value = phone
                            confidence = 1.0

                            if src.get("social_platform"):
                                display = " ".join(x for x in [n, f] if x).strip() or None
                                _set_social(
                                    conn, cid, src["social_platform"],
                                    exists_status="yes",
                                    username=r["username"] if "username" in r.keys() else None,
                                    display_name=display,
                                    user_id=str(r["user_id"]) if "user_id" in r.keys() and r["user_id"] else None,
                                    raw_data=r["raw_data"] if "raw_data" in r.keys() else None,
                                )

                            _add_link(conn, cid, f"{src['name']}_db",
                                      src["table_name"], r[src["pk_column"]],
                                      match_key, match_value, confidence)
                            stats[match_key] = stats.get(match_key, 0) + 1
                            matched += 1
                            continue

                    stats["no_match"] = stats.get("no_match", 0) + 1
                    continue

                stats[match_key] = stats.get(match_key, 0) + 1

                n = r[first_col] if first_col else None
                f = r[last_col] if last_col else None
                if n or f:
                    _update_contact_names(conn, cid, n, f)

                if nc_col:
                    nc = r[nc_col]
                    if nc and not normalize_phone(nc):
                        _update_contact_national_code(conn, cid, nc)

                if src.get("social_platform"):
                    display = " ".join(x for x in [n, f] if x).strip() or None
                    _set_social(
                        conn, cid, src["social_platform"],
                        exists_status="yes",
                        username=r["username"] if "username" in r.keys() else None,
                        display_name=display,
                        user_id=str(r["user_id"]) if "user_id" in r.keys() and r["user_id"] else None,
                        raw_data=r["raw_data"] if "raw_data" in r.keys() else None,
                    )

                _add_link(conn, cid, f"{src['name']}_db",
                          src["table_name"], r[src["pk_column"]],
                          match_key, match_value, confidence)
                matched += 1
    finally:
        src_conn.close()

    invalidate_external_cache(src["name"])
    log(f"[sync] {src['name']}: {matched} rows matched.")
    for k, v in stats.items():
        log(f"        {k:24s} -> {v}")
    return matched


def run_sync():
    log("[sync] starting ...")
    total = 0
    for src in list_data_sources(only_enabled=True):
        try:
            total += match_with_source(src)
        except Exception as e:
            log(f"[sync] source '{src['name']}' failed: {e}")
    audit("sync", f"total={total}")
    log("[sync] done.")
    return total


# ------------------------------------------------------------
# Query
# ------------------------------------------------------------
def query_contacts(search="", status="all", platform="all", source="all"):
    sql = """
        SELECT
            c.id,
            c.phone_number,
            c.status,
            c.name,
            c.last_name,
            c.national_code,
            c.notes,
            c.photo_path,
            c.pinned,
            c.special,
            lsrc.srcs        AS linked_sources
        FROM contacts c
        LEFT JOIN (
            SELECT contact_id,
                   GROUP_CONCAT(DISTINCT source_db) AS srcs
            FROM contact_links
            GROUP BY contact_id
        ) lsrc ON lsrc.contact_id = c.id
    """
    where, params = [], []

    if search:
        like = f"%{search}%"
        where.append("""
            (c.phone_number LIKE ?
             OR c.name LIKE ?
             OR c.last_name LIKE ?
             OR c.national_code LIKE ?
             OR c.notes LIKE ?
             OR EXISTS (
                 SELECT 1 FROM contact_socials s
                 WHERE s.contact_id = c.id
                   AND (s.username LIKE ? OR s.user_id LIKE ?)
             ))
        """)
        params += [like] * 7

    if status and status != "all":
        where.append("c.status = ?")
        params.append(status)

    if platform == "has_identity":
        where.append("EXISTS (SELECT 1 FROM contact_links l "
                     "WHERE l.contact_id=c.id AND l.source_db='identity_db')")
    elif platform == "no_identity":
        where.append("NOT EXISTS (SELECT 1 FROM contact_links l "
                     "WHERE l.contact_id=c.id AND l.source_db='identity_db')")
    elif platform == "pinned":
        where.append("c.pinned = 1")
    elif platform == "special":
        where.append("c.special = 1")
    elif platform == "with_photo":
        where.append("EXISTS (SELECT 1 FROM contact_images i "
                     "WHERE i.contact_id = c.id)")
    elif platform == "no_photo":
        where.append("NOT EXISTS (SELECT 1 FROM contact_images i "
                     "WHERE i.contact_id = c.id)")
    elif platform != "all":
        if platform.endswith("_yes"):
            p = platform[:-4]
            where.append(
                "EXISTS (SELECT 1 FROM contact_socials s "
                "WHERE s.contact_id=c.id AND s.platform=? "
                "AND s.exists_status='yes')"
            )
            params.append(p)
        elif platform.endswith("_no"):
            p = platform[:-3]
            where.append(
                "EXISTS (SELECT 1 FROM contact_socials s "
                "WHERE s.contact_id=c.id AND s.platform=? "
                "AND s.exists_status='no')"
            )
            params.append(p)
        elif platform.endswith("_unknown"):
            p = platform[:-8]
            where.append(
                "EXISTS (SELECT 1 FROM contact_socials s "
                "WHERE s.contact_id=c.id AND s.platform=? "
                "AND s.exists_status='unknown')"
            )
            params.append(p)
        else:
            where.append(
                "EXISTS (SELECT 1 FROM contact_socials s "
                "WHERE s.contact_id=c.id AND s.platform=?)"
            )
            params.append(platform)

    if source == "only_contacts":
        where.append("NOT EXISTS (SELECT 1 FROM contact_links l "
                     "WHERE l.contact_id = c.id)")
    elif source.startswith("with_"):
        sname = source[5:]
        where.append("EXISTS (SELECT 1 FROM contact_links l "
                     "WHERE l.contact_id=c.id AND l.source_db=?)")
        params.append(f"{sname}_db")

    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY c.pinned DESC, c.special DESC, c.id ASC"

    with connect(DB_PATH) as conn:
        socials_by_cid = {}
        for s in conn.execute(
            "SELECT contact_id, platform, exists_status, username, user_id "
            "FROM contact_socials"
        ):
            socials_by_cid.setdefault(s["contact_id"], {})[s["platform"]] = dict(s)

        has_img_rows = conn.execute(
            "SELECT DISTINCT contact_id FROM contact_images"
        ).fetchall()
        has_image = {r["contact_id"] for r in has_img_rows}

        labels = source_label_map()
        out = []
        for r in conn.execute(sql, params):
            d = dict(r)
            codes = ["contacts"]
            linked = (d.pop("linked_sources") or "").split(",")
            for db_name in linked:
                if db_name and db_name != "contacts" and db_name not in codes:
                    codes.append(db_name)
            d["sources"] = " ".join(labels.get(c, c) for c in codes)
            d["socials"] = socials_by_cid.get(d["id"], {})
            d["has_image"] = d["id"] in has_image
            out.append(d)
    return out


def contact_basic(contact_id):
    with connect(DB_PATH) as conn:
        c = conn.execute(
            "SELECT * FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
        if not c:
            return None
        socials = [dict(r) for r in conn.execute(
            "SELECT * FROM contact_socials WHERE contact_id=? ORDER BY platform",
            (contact_id,)
        )]
        links = [dict(r) for r in conn.execute(
            "SELECT * FROM contact_links WHERE contact_id=? ORDER BY source_db, id",
            (contact_id,)
        )]
        rows = conn.execute(
            "SELECT DISTINCT source_db FROM contact_links WHERE contact_id=?",
            (contact_id,)
        ).fetchall()

    linked = {r["source_db"] for r in rows}
    codes = ["contacts"]
    for db_name in sorted(linked):
        if db_name not in codes:
            codes.append(db_name)
    return {
        "contact": dict(c),
        "socials": socials,
        "links":   links,
        "sources": codes,
    }


def get_contact_detail(contact_id):
    out = {"contact": {}, "socials": [], "identity": {}, "telegram": {},
           "links": [], "sources": [], "external": {}, "files": [],
           "images": [], "primary_image_id": None}

    with connect(DB_PATH) as conn:
        c = conn.execute(
            "SELECT * FROM contacts WHERE id=?", (contact_id,)
        ).fetchone()
        if not c:
            return out
        out["contact"] = dict(c)

        for s in conn.execute(
            "SELECT * FROM contact_socials WHERE contact_id=? ORDER BY platform",
            (contact_id,)
        ):
            out["socials"].append(dict(s))

        for l in conn.execute(
            "SELECT * FROM contact_links WHERE contact_id=? ORDER BY source_db, id",
            (contact_id,)
        ):
            out["links"].append(dict(l))

        rows = conn.execute(
            "SELECT DISTINCT source_db FROM contact_links WHERE contact_id=?",
            (contact_id,)
        ).fetchall()

        out["images"] = [dict(r) for r in conn.execute("""
            SELECT id, contact_id, name, mime, size_bytes, width, height,
                   sha256, is_primary, encoding, created_at
            FROM contact_images
            WHERE contact_id=?
            ORDER BY is_primary DESC, id ASC
        """, (contact_id,))]

    linked = {r["source_db"] for r in rows}
    codes = ["contacts"]
    for db_name in sorted(linked):
        if db_name not in codes:
            codes.append(db_name)
    out["sources"] = codes

    if out["images"]:
        out["primary_image_id"] = out["images"][0]["id"]

    phone = out["contact"].get("phone_number")
    external = {}
    if phone:
        for src in list_data_sources(only_enabled=True):
            idx = _get_index(src)
            row = idx.get(phone)
            if row:
                external[src["name"]] = row
    out["external"] = external
    out["identity"] = external.get("identity", {})
    out["telegram"] = external.get("telegram", {})

    out["files"] = list_contact_files(contact_id)
    return out


# ------------------------------------------------------------
# Export
# ------------------------------------------------------------
def build_report_rows():
    plats = platform_names()
    pieces = []
    for p in plats:
        pieces.append(
            f"MAX(CASE WHEN s.platform='{p}' THEN s.exists_status END) "
            f'AS "{p}"'
        )
        pieces.append(
            f"MAX(CASE WHEN s.platform='{p}' THEN s.username END) "
            f'AS "{p} username"'
        )
        pieces.append(
            f"MAX(CASE WHEN s.platform='{p}' THEN s.display_name END) "
            f'AS "{p} display name"'
        )
        pieces.append(
            f"MAX(CASE WHEN s.platform='{p}' THEN s.user_id END) "
            f'AS "{p} userid"'
        )
    body = ",\n            ".join(pieces)

    sql = f"""
        SELECT
            c.id,
            c.phone_number, c.status, c.name, c.last_name, c.national_code,
            c.notes, c.photo_path, c.pinned, c.special,
            {body},
            (SELECT GROUP_CONCAT(DISTINCT source_db)
             FROM contact_links
             WHERE contact_id = c.id) AS linked_sources,
            (SELECT COUNT(*) FROM contact_images
             WHERE contact_id = c.id) AS image_count
        FROM contacts c
        LEFT JOIN contact_socials s ON s.contact_id = c.id
        GROUP BY c.id
        ORDER BY c.pinned DESC, c.special DESC, c.id ASC
    """

    labels = source_label_map()
    out = []
    with connect(DB_PATH) as conn:
        for r in conn.execute(sql):
            d = dict(r)
            linked = (d.pop("linked_sources") or "").split(",")
            codes = ["contacts"]
            for db_name in linked:
                if db_name and db_name != "contacts" and db_name not in codes:
                    codes.append(db_name)
            d["sources"] = " ".join(labels.get(c, c) for c in codes)
            out.append(d)
    return out


def export_csv(path="reports/contacts_wide.csv"):
    rows = build_report_rows()
    if not rows:
        log("[export] no data.")
        return None
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    log(f"[export] CSV -> {path}")
    audit("export_csv", path)
    return path


def export_excel(path="reports/contacts_wide.xlsx"):
    try:
        from openpyxl import Workbook
    except ImportError:
        log("[export] openpyxl not installed. Run: pip install openpyxl")
        return None
    rows = build_report_rows()
    if not rows:
        log("[export] no data.")
        return None
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("contacts")
    ws.title = "contacts"
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r[h] for h in headers])
    wb.save(path)
    log(f"[export] Excel -> {path}")
    audit("export_excel", path)
    return path